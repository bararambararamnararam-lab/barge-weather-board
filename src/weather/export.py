"""표를 엑셀 파일(.xlsx)로 내보내는 부분.

`9_엑셀로_내보내기.bat` 을 더블클릭하면 실행된다.
결과 파일은 claude-weather\\data\\export 폴더에 날짜 이름으로 저장된다.

만들어지는 파일 구조:
  시트 하나 = 탭 하나 (풍속·풍향, 유의파고, 운항 판단 ...)
  행 = 위치, 열 = 시각, 셀 색 = 초록/노랑/빨강/회색
  셀 메모(마우스 올리면 뜨는 설명) = 화면 툴팁과 같은 내용
"""

from __future__ import annotations

import sys
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import judge
from .config import PROJECT_ROOT, Config, load_config
from .db import COLUMN_LABELS, VALUE_COLUMNS, fetch_forecast, fetch_latest_warnings, session

EXPORT_DIR = PROJECT_ROOT / "data" / "export"

# 시트 이름과 그 시트에 넣을 열
SHEETS: list[tuple[str, list[str]]] = [
    ("풍속·풍향", ["wind_speed_ms", "wind_direction_deg"]),
    ("순간풍속", ["wind_gust_ms"]),
    ("유의파고", ["wave_height_m"]),
    ("파향·파주기", ["wave_direction_deg", "wave_period_s"]),
    ("시정", ["visibility_km"]),
    ("강수", ["precipitation_mm"]),
]

FILLS = {
    judge.NORMAL: PatternFill("solid", fgColor="E6F4EA"),
    judge.CAUTION: PatternFill("solid", fgColor="FFF4D6"),
    judge.UNAVAILABLE: PatternFill("solid", fgColor="FDE7E7"),
    judge.NO_DATA: PatternFill("solid", fgColor="F0F0F0"),
}

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")


def _time_columns(rows: list[dict], step_hours: int) -> list[str]:
    times = sorted({r["valid_time"] for r in rows})
    if step_hours <= 1:
        return times
    return [t for t in times if int(t[11:13]) % step_hours == 0]


def _write_sheet(
    workbook: Workbook,
    title: str,
    config: Config,
    location_ids: list[str],
    rows: list[dict],
    times: list[str],
    columns: list[str],
    warning_rows: list[dict],
    collected_at: str | None,
    judgement_mode: bool = False,
) -> None:
    sheet = workbook.create_sheet(title=title[:31])
    indexed = {(r["location_id"], r["valid_time"]): r for r in rows}

    sheet.cell(row=1, column=1, value="위치").font = Font(bold=True)
    sheet.cell(row=1, column=1).border = BORDER
    for index, valid_time in enumerate(times, start=2):
        cell = sheet.cell(row=1, column=index,
                          value="{} {}".format(valid_time[5:10], valid_time[11:16]))
        cell.font = Font(bold=True, size=9)
        cell.alignment = CENTER
        cell.border = BORDER

    for row_index, location_id in enumerate(location_ids, start=2):
        location = config.locations[location_id]
        location_warnings = judge.warnings_for_location(config, warning_rows, location_id)

        name_cell = sheet.cell(row=row_index, column=1, value=location.name)
        name_cell.font = Font(bold=True)
        name_cell.border = BORDER

        for column_index, valid_time in enumerate(times, start=2):
            record = indexed.get((location_id, valid_time))
            cell = sheet.cell(row=row_index, column=column_index)
            cell.border = BORDER
            cell.alignment = CENTER

            if record is None:
                cell.value = "-"
                cell.fill = FILLS[judge.NO_DATA]
                continue

            hits = judge.active_warnings_at(config, location_warnings, valid_time)
            verdict = judge.judge_cell(config, record, hits)
            tooltip = judge.build_tooltip(
                verdict, location.name, valid_time,
                record.get("collected_at") or collected_at, "예보")

            if judgement_mode:
                status = verdict.status
                cell.value = judge.JUDGEMENT_LABELS_KO[status]
            else:
                status = judge.NORMAL
                pieces = []
                any_value = False
                for column in columns:
                    value = record.get(column)
                    if value is not None:
                        any_value = True
                    pieces.append(judge.format_value(column, value))
                    metric = judge.evaluate_metric(column, value, config.thresholds)
                    if metric.status == judge.NO_DATA:
                        continue
                    if judge.SEVERITY[metric.status] > judge.SEVERITY[status]:
                        status = metric.status
                if not any_value:
                    status = judge.NO_DATA
                    cell.value = "-"
                else:
                    cell.value = " / ".join(p for p in pieces if p)

            cell.fill = FILLS[status]
            comment = Comment(tooltip, "기상조회")
            comment.width = 320
            comment.height = 240
            cell.comment = comment

    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 26
    for index in range(2, len(times) + 2):
        sheet.column_dimensions[get_column_letter(index)].width = 13


def export_route(
    route_id: str,
    days_ahead: int = 7,
    step_hours: int = 3,
    config: Config | None = None,
) -> Path:
    """한 항로를 엑셀 파일 하나로 내보낸다. 만들어진 파일 경로를 돌려준다."""
    cfg = config or load_config()
    locations = cfg.route_locations(route_id)
    location_ids = [loc.id for loc in locations]

    now = datetime.now(cfg.timezone).replace(minute=0, second=0, microsecond=0)
    valid_from = now.strftime("%Y-%m-%dT%H:%M")
    valid_to = (now + timedelta(days=days_ahead)).strftime("%Y-%m-%dT%H:%M")

    with session() as conn:
        rows = [dict(r) for r in fetch_forecast(conn, location_ids, None,
                                                valid_from, valid_to)]
        warning_collected, warning_rows = fetch_latest_warnings(conn)
        warning_rows = [dict(r) for r in warning_rows]

    if not rows:
        raise RuntimeError(
            "내보낼 예보 데이터가 없습니다. 1_수집하기.bat 을 먼저 실행하세요."
        )

    times = _time_columns(rows, step_hours)
    collected_at = max(r["collected_at"] for r in rows)

    workbook = Workbook()
    workbook.remove(workbook.active)

    # 표지 시트
    cover = workbook.create_sheet("안내")
    lines = [
        ["한·중 바지선 운항 기상 조회"],
        [],
        ["항로", "{} · {}".format(cfg.routes[route_id].label,
                                 cfg.routes[route_id].name)],
        ["작성 시각", now.strftime("%Y-%m-%d %H:%M") + " KST"],
        ["예보 수집 시각", collected_at.replace("T", " ") + " KST"],
        ["특보 수집 시각", (warning_collected or "-").replace("T", " ")],
        ["표시 기간", "{} ~ {}".format(valid_from.replace("T", " "),
                                        valid_to.replace("T", " "))],
        ["표시 간격", "{}시간".format(step_hours)],
        [],
        ["색상", "초록=정상, 노랑=주의, 빨강=불가, 회색=데이터 없음"],
        ["툴팁", "셀에 마우스를 올리면 판정 근거가 메모로 나옵니다."],
        [],
        ["주의", "이 판정은 화면 경보용입니다. 자동 출항 승인 기준이 아닙니다."],
        ["", "실제 피항·출항 판단은 담당자가 하십시오."],
        [],
        ["출처", "ECMWF IFS+WAM (CC BY 4.0), NOAA GFS, MET Norway (CC BY 4.0), 기상청 API 허브"],
        ["라이선스", "ECMWF·MET Norway 는 CC BY 4.0, NOAA 는 미국 공공저작물, 기상청은 공공데이터. 모두 상업적 사용이 허용됩니다."],
        ["가공 사실", "원자료를 그대로 옮긴 것이 아니라 운항 판단을 계산해 넣은 가공 자료입니다. (CC BY 4.0 은 변경 사실 표시를 요구합니다)"],
        ["해양 예보 한계", "파고·파주기는 약 9~10일까지만 값이 있고 그 이후는 회색입니다."],
    ]
    for line in lines:
        cover.append(line)
    cover.column_dimensions["A"].width = 18
    cover.column_dimensions["B"].width = 78
    for row in cover.iter_rows(min_col=1, max_col=1):
        for cell in row:
            cell.font = Font(bold=True)

    for title, columns in SHEETS:
        _write_sheet(workbook, title, cfg, location_ids, rows, times, columns,
                     warning_rows, collected_at)

    _write_sheet(workbook, "운항 판단", cfg, location_ids, rows, times, [],
                 warning_rows, collected_at, judgement_mode=True)

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = "기상_{}_{}.xlsx".format(
        cfg.routes[route_id].label, now.strftime("%Y%m%d_%H%M"))
    output_path = EXPORT_DIR / filename
    workbook.save(output_path)
    return output_path


def main(argv: list[str] | None = None) -> int:
    args = list(argv if argv is not None else sys.argv[1:])
    try:
        config = load_config()
    except Exception as exc:
        print("설정 파일을 읽지 못했습니다.\n{}".format(exc))
        return 1

    route_ids = [a.upper() for a in args if a.upper() in config.routes]
    if not route_ids:
        route_ids = list(config.routes.keys())

    print("엑셀 파일을 만듭니다. 대상 항로: {}".format(
        ", ".join(config.routes[r].label for r in route_ids)))
    failed = False
    for route_id in route_ids:
        try:
            path = export_route(route_id, config=config)
            print("  완료: {}".format(path))
        except Exception as exc:
            failed = True
            print("  실패({}): {}".format(route_id, exc))

    if not failed:
        print("\n저장 위치: {}".format(EXPORT_DIR))
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
